import pandas as pd
import os
import shutil
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter import scrolledtext
import threading
import queue
import re
import datetime as dt
import numpy as np
import traceback

# --- 설정값 ---
TARGET_SLDPRT_EXTENSION = ".sldprt"
TARGET_JPG_EXTENSION = ".jpg"
SLDPRT_EXCLUSION_FILE_PATH = r'\\sechang-nas\세창산업\연구소\Reflow 파트\진행 중\설계실 자료\CSENG_DATA.xlsx'
TOTAL_BOM_PATH = r'\\sechang-nas\세창산업\연구소\Reflow 파트\진행 중\설계실 자료\TOTAL BOM.xls'
ECN_STATUS_PATH = r'\\sechang-nas\세창산업\연구소\Reflow 파트\진행 중\설계실 자료\연구소2팀\ECN_현황.xlsx'
JPG_SOURCE_PATH = r'\\sechang-nas\세창산업\연구소\Reflow 파트\4 도면'

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("Advanced 도면 파일 분류기 v6.6 (Variable Fix)")
        self.root.minsize(700, 520)

        self.source_folder = tk.StringVar()
        self.main_excel_file = tk.StringVar()
        self.status_text = tk.StringVar(value="대기 중...")
        self.copy_jpg_var = tk.BooleanVar(value=True)

        # --- GUI ---
        main_frame = ttk.Frame(root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        main_frame.columnconfigure(1, weight=1)
        for r in (0, 1, 2, 3, 4, 5):
            main_frame.rowconfigure(r, weight=0)
        main_frame.rowconfigure(6, weight=1)

        path_frame = ttk.LabelFrame(main_frame, text="경로 설정 및 옵션", padding="10")
        path_frame.grid(row=0, column=0, columnspan=3, sticky="ew")
        path_frame.columnconfigure(1, weight=1)

        ttk.Label(path_frame, text="1. SLDPRT 원본:").grid(row=0, column=0, padx=5, pady=8, sticky="w")
        ttk.Entry(path_frame, textvariable=self.source_folder, state="readonly").grid(row=0, column=1, sticky="ew")
        self.source_button = ttk.Button(path_frame, text="폴더 선택", command=self.select_source_folder)
        self.source_button.grid(row=0, column=2, padx=5)

        ttk.Label(path_frame, text="2. 엑셀 파일:").grid(row=1, column=0, padx=5, pady=8, sticky="w")
        ttk.Entry(path_frame, textvariable=self.main_excel_file, state="readonly").grid(row=1, column=1, sticky="ew")
        self.excel_button = ttk.Button(path_frame, text="파일 선택", command=self.select_main_excel)
        self.excel_button.grid(row=1, column=2, padx=5)

        self.jpg_checkbutton = ttk.Checkbutton(path_frame, text="JPG 파일 추가 복사 (고정 경로 사용)", variable=self.copy_jpg_var)
        self.jpg_checkbutton.grid(row=2, column=0, columnspan=3, padx=5, pady=8, sticky="w")

        ttk.Label(main_frame, textvariable=self.status_text, font=("", 9)).grid(row=2, column=0, columnspan=3, sticky="w", padx=5, pady=(10,0))
        self.progress_bar = ttk.Progressbar(main_frame, orient='horizontal', mode='determinate')
        self.progress_bar.grid(row=3, column=0, columnspan=3, sticky="ew", padx=5, pady=5)
        self.start_button = ttk.Button(main_frame, text="🚀 작업 시작", command=self.start_task_thread, state="disabled")
        self.start_button.grid(row=4, column=0, columnspan=3, pady=10, sticky="ew")

        log_frame = ttk.LabelFrame(main_frame, text="작업 로그", padding="8")
        log_frame.grid(row=6, column=0, columnspan=3, sticky="nsew")
        self.log_text = scrolledtext.ScrolledText(log_frame, height=12, state="disabled", font=("", 9))
        self.log_text.pack(fill=tk.BOTH, expand=True)

    def _append_log_to_gui(self, message: str, level: str = "INFO"):
        timestamp = dt.datetime.now().strftime("%H:%M:%S")
        self.log_text.config(state="normal")
        self.log_text.insert(tk.END, f"[{timestamp}] [{level}] {message}\n")
        self.log_text.see(tk.END)
        self.log_text.config(state="disabled")

    def _set_controls_enabled(self, enabled: bool):
        state = "normal" if enabled else "disabled"
        for w in [self.start_button, self.source_button, self.excel_button, self.jpg_checkbutton]:
            w.config(state=state)

    def select_source_folder(self):
        folder = filedialog.askdirectory(title="1. SLDPRT 파일이 있는 '원본 폴더'를 선택하세요")
        if folder: self.source_folder.set(folder)
        self.check_paths()

    def select_main_excel(self):
        file = filedialog.askopenfilename(title="2. 필요한 모든 열이 포함된 '메인 엑셀 파일'을 선택하세요", filetypes=[("Excel files", "*.xlsx *.xls")])
        if file: self.main_excel_file.set(file)
        self.check_paths()

    def check_paths(self):
        if self.source_folder.get() and self.main_excel_file.get():
            self.start_button.config(state="normal")

    def start_task_thread(self):
        self._set_controls_enabled(False)
        self.log_text.config(state="normal")
        self.log_text.delete(1.0, tk.END)
        self.log_text.config(state="disabled")
        self.task_queue = queue.Queue()
        threading.Thread(target=self.run_processing_task, args=(self.task_queue,), daemon=True).start()
        self.root.after(100, self.process_queue)

    def process_queue(self):
        try:
            message = self.task_queue.get_nowait()
        except queue.Empty:
            self.root.after(100, self.process_queue)
            return

        if isinstance(message, dict):
            mtype = message.get('type')
            if mtype == 'progress':
                self.progress_bar['value'] = message.get('progress', 0)
                self.status_text.set(message.get('status', ''))
            elif mtype == 'gui_log':
                self._append_log_to_gui(message['text'])
            elif mtype == 'error':
                self._append_log_to_gui(message.get('log_text', '알 수 없는 오류'), level="ERROR")
                messagebox.showerror("오류 발생", message.get('popup_text', '알 수 없는 오류'))
                self.status_text.set("오류 발생. 다시 시도하세요.")
                self.progress_bar['value'] = 0
                self._set_controls_enabled(True)
            elif mtype == 'done':
                self.progress_bar['value'] = 100
                self.status_text.set(message.get('status', '완료'))
                self._append_log_to_gui("모든 작업이 완료되었습니다.")
                if message.get('log_save_path'):
                    self._append_log_to_gui(f"상세 로그가 엑셀 파일로 저장되었습니다: {message['log_save_path']}")
                self._set_controls_enabled(True)
                messagebox.showinfo("작업 완료", message.get('final_message', '작업이 완료되었습니다.'))
        
        self.root.after(100, self.process_queue)

    def run_processing_task(self, q: queue.Queue):
        excel_log_records = []
        def log_for_excel(message):
            excel_log_records.append(message)
        
        def emit_progress(pct, status, gui_log=None):
            q.put({'type': 'progress', 'progress': pct, 'status': status})
            if gui_log:
                q.put({'type': 'gui_log', 'text': gui_log})

        def emit_error(popup_text, log_text):
            q.put({'type': 'error', 'popup_text': popup_text, 'log_text': log_text})
        
        def emit_done(status, final_message, log_save_path=None):
            q.put({'type': 'done', 'status': status, 'final_message': final_message, 'log_save_path': log_save_path})

        try:
            source = self.source_folder.get()
            excel_path = self.main_excel_file.get()
            copy_jpg_enabled = self.copy_jpg_var.get()
            
            parent_dest_folder = os.path.join(source, "Export")
            os.makedirs(parent_dest_folder, exist_ok=True)
            
            sldprt_dest_folder = os.path.join(parent_dest_folder, "SLDPRT_Export")
            jpg_dest_folder = os.path.join(parent_dest_folder, "JPG_Export")
            os.makedirs(sldprt_dest_folder, exist_ok=True)
            
            report_excel_path = os.path.join(parent_dest_folder, "LIST.xlsx")
            ts_str = dt.datetime.now().strftime('%Y%m%d_%H%M%S')
            log_excel_path = os.path.join(parent_dest_folder, f"RUN_LOG_{ts_str}.xlsx")

            # --- 엑셀 사전 처리 단계 ---
            emit_progress(5, "엑셀 파일 사전 처리 중...")
            try:
                from openpyxl import load_workbook
                from openpyxl.utils import get_column_letter, column_index_from_string
                from openpyxl.styles import Font
                from io import BytesIO

                with open(excel_path, "rb") as f:
                    in_mem_file = BytesIO(f.read())
                
                wb = load_workbook(in_mem_file)
                ws = wb.active

                if ws.merged_cells.ranges:
                    for merged_range in list(ws.merged_cells.ranges):
                        ws.unmerge_cells(str(merged_range))

                rows_to_delete = []
                target_column = column_index_from_string('D')
                for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
                    cell = row[target_column - 1]
                    color = cell.font.color
                    if color is not None and color.type == "rgb" and color.rgb and color.rgb.upper() == "FFFF0000":
                        rows_to_delete.append(cell.row)
                    elif color is not None and color.type == "theme":
                        theme_colors = {
                            0: "FF000000",
                            1: "FFFFFFFF",
                            2: "FFFF0000",
                            3: "FF00FF00",
                            4: "FF0000FF",
                        }
                        rgb = theme_colors.get(color.theme)
                        if rgb == "FFFF0000":
                            rows_to_delete.append(cell.row)

                for offset, row_idx in enumerate(sorted(set(rows_to_delete))):
                    ws.delete_rows(row_idx - offset)

                rows_deleted = 0
                if ws['A8'].value == 'Peer Review by' and ws['A9'].value == 'Peer Review Date':
                    ws.delete_rows(8, 2)
                    rows_deleted = 2

                header_row_index = 12 - rows_deleted
                header_excel_row = header_row_index + 1

                expected_headers = ['E', 'F', 'G', 'H', 'I', 'J', 'K']
                existing_headers = [ws.cell(row=header_excel_row, column=5 + idx).value for idx in range(len(expected_headers))]
                if existing_headers != expected_headers:
                    ws.insert_cols(idx=5, amount=len(expected_headers))
                
                temp_excel_file = BytesIO()
                wb.save(temp_excel_file)
                temp_excel_file.seek(0)
                
                df_raw = pd.read_excel(temp_excel_file, header=header_row_index)
                
                required_cols = ['PN', 'DRAWING No.', 'P-QTY', 'QTY', 'KITID']
                missing_cols = [col for col in required_cols if col not in df_raw.columns]
                if missing_cols:
                    excel_header_row = header_row_index + 1
                    raise ValueError(f"엑셀 파일의 헤더({excel_header_row}행)에 필요한 열이 없습니다: {', '.join(missing_cols)}")
                
                df_filtered = df_raw[df_raw['KITID'].isin(['KITM', 'SDRWM', '9DRWM'])].copy()

                d_series = df_filtered['PN'].fillna("").astype(str)
                d_series = d_series.replace({'nan': ''})
                d_lengths = d_series.str.len()
                df_filtered.loc[:, 'E_new'] = d_lengths

                first_chars = d_series.str[0].fillna('')
                f_mapping = {'1': 1, '2': 2, '6': 5, '5': 5, '8': 4, '4': 4}
                f_values = first_chars.map(f_mapping)
                f_values = f_values.where(d_lengths > 2, pd.NA)
                df_filtered.loc[:, 'F_new'] = pd.Series(pd.array(f_values, dtype=pd.Int64Dtype()), index=df_filtered.index)

                allowed_for_g = {'1', '2', '4', '5', '6', '8'}
                core_slice = d_series.str.slice(1, 9)
                df_filtered.loc[:, 'G_new'] = np.where((d_lengths > 2) & first_chars.isin(allowed_for_g), core_slice, "")

                f_component = df_filtered['F_new'].apply(lambda v: "" if pd.isna(v) else str(int(v)))
                g_component = df_filtered['G_new'].fillna("")
                df_filtered.loc[:, 'H_new'] = f_component + g_component
                df_filtered.loc[:, 'I_new'] = pd.to_numeric(df_filtered['H_new'], errors='coerce')
                df_filtered.loc[:, 'J_new'] = pd.to_numeric(df_filtered['I_new'], errors='coerce')

                k_cond = df_filtered['E_new'].isin([7, 9])
                df_filtered.loc[:, 'K_new'] = np.where(k_cond, pd.to_numeric(df_filtered['H_new'].str.slice(0, 7), errors='coerce'), pd.to_numeric(df_filtered['H_new'].str.slice(0, 6), errors='coerce'))

                df_total_bom = pd.read_excel(TOTAL_BOM_PATH, sheet_name='standard', usecols="B", dtype=str)
                df_ecn = pd.read_excel(ECN_STATUS_PATH, sheet_name='Sheet1', usecols="D,E", dtype=str)

                total_bom_set = set(df_total_bom.iloc[:, 0].dropna())
                ecn_map_z = df_ecn.set_index(df_ecn.columns[0])[df_ecn.columns[1]].to_dict()

                j_series = df_filtered['J_new'].dropna().astype('Int64').astype(str)
                df_filtered.loc[:, 'X_new'] = j_series.apply(lambda x: x if x in total_bom_set else np.nan)
                df_filtered.loc[:, 'Y_new'] = j_series.apply(lambda x: x if x in ecn_map_z else np.nan)
                df_filtered.loc[:, 'Z_new'] = j_series.map(ecn_map_z)

                new_cols_df = df_filtered[['E_new', 'F_new', 'G_new', 'H_new', 'I_new', 'J_new', 'K_new', 'X_new', 'Y_new', 'Z_new']]
                df_processed = df_raw.merge(new_cols_df, left_index=True, right_index=True, how='left')

                original_cols = df_raw.columns.tolist()
                new_col_names = ['E', 'F', 'G', 'H', 'I', 'J', 'K']

                try:
                    pn_index = original_cols.index('PN')
                    final_cols_order = original_cols[:pn_index+1] + new_col_names + original_cols[pn_index+1:]
                except ValueError:
                    final_cols_order = original_cols + new_col_names

                rename_map_new = {f'{c}_new': c for c in ['E', 'F', 'G', 'H', 'I', 'J', 'K', 'X', 'Y', 'Z']}
                df_processed.rename(columns=rename_map_new, inplace=True)

                all_possible_cols = list(df_processed.columns)
                final_cols_order_full = [col for col in (final_cols_order + ['X', 'Y', 'Z']) if col in all_possible_cols]
                df_processed_final = df_processed.reindex(columns=final_cols_order_full)

                filtered_indices = df_filtered.index.to_numpy()
                filtered_index_set = set(filtered_indices)

                source_indices = df_processed_final.index.to_numpy()
                alpha_mask = df_processed_final['PN'].astype(str).str.contains('[A-Za-z]', na=False)
                if alpha_mask.any():
                    keep_mask = ~alpha_mask.to_numpy()
                    df_processed_final = df_processed_final.loc[keep_mask].copy()
                    source_indices = source_indices[keep_mask]

                filtered_flags = np.isin(source_indices, list(filtered_index_set))
                if hasattr(filtered_flags, 'tolist'):
                    filtered_flags = filtered_flags.tolist()

                df_processed_final.reset_index(drop=True, inplace=True)

                processed_filename = os.path.splitext(os.path.basename(excel_path))[0] + "_판금리스트.xlsx"
                processed_filepath = os.path.join(os.path.dirname(excel_path), processed_filename)

                header_excel_row = header_row_index + 1
                data_start_row = header_row_index + 2
                total_data_rows = len(df_processed_final)
                last_data_row = data_start_row + total_data_rows - 1 if total_data_rows else data_start_row - 1

                target_columns = ['E', 'F', 'G', 'H', 'I', 'J', 'K', 'X', 'Y', 'Z']
                blue_bold_font = Font(color="FF0000FF", bold=True)
                highlight_columns = {'X', 'Y', 'Z'}
                for col_letter in target_columns:
                    cell = ws[f"{col_letter}{header_excel_row}"]
                    cell.value = col_letter
                    if col_letter in highlight_columns:
                        cell.font = blue_bold_font

                if total_data_rows:
                    for row_num in range(data_start_row, last_data_row + 1):
                        for col_letter in target_columns:
                            ws[f"{col_letter}{row_num}"] = None

                    total_rows = len(df_processed_final)
                    for pos in range(total_rows):
                        row_num = data_start_row + pos
                        if not (data_start_row <= row_num <= last_data_row):
                            continue

                        row_series = df_processed_final.iloc[pos]

                        for col_name, col_letter in zip(['E', 'F', 'G', 'H', 'I', 'J', 'K', 'X', 'Y', 'Z'], ['E', 'F', 'G', 'H', 'I', 'J', 'K', 'X', 'Y', 'Z']):
                            value = row_series.get(col_name)
                            if pd.isna(value) or value == "":
                                ws[f"{col_letter}{row_num}"] = None
                            else:
                                ws[f"{col_letter}{row_num}"] = value

                        if filtered_flags[pos]:
                            for col_letter in highlight_columns:
                                ws[f"{col_letter}{row_num}"].font = blue_bold_font
                            ws.row_dimensions[row_num].hidden = False
                        else:
                            ws.row_dimensions[row_num].hidden = True

                    rows_to_delete_d = []
                    for row_idx in range(data_start_row, last_data_row + 1):
                        cell_value = ws[f"D{row_idx}"].value
                        if isinstance(cell_value, str) and any(ch.isalpha() for ch in cell_value):
                            rows_to_delete_d.append(row_idx)

                    for row_idx in reversed(rows_to_delete_d):
                        ws.delete_rows(row_idx)

                    remaining_rows = max(total_rows - len(rows_to_delete_d), 0)
                    first_extra_row = data_start_row + remaining_rows
                    current_max_row = ws.max_row
                    extra_rows = current_max_row - (first_extra_row - 1)
                    if extra_rows > 0:
                        ws.delete_rows(first_extra_row, extra_rows)

                    if remaining_rows:
                        last_data_row = first_extra_row - 1
                        total_data_rows = remaining_rows
                    else:
                        last_data_row = data_start_row - 1
                        total_data_rows = 0
                else:
                    current_max_row = ws.max_row
                    extra_rows = current_max_row - (data_start_row - 1)
                    if extra_rows > 0:
                        ws.delete_rows(data_start_row, extra_rows)

                df_col_letter_map = {get_column_letter(idx): col_name for idx, col_name in enumerate(df_processed_final.columns, start=1)}

                def _value_length(value):
                    if value is None:
                        return 0
                    text = str(value)
                    return len(text.strip())

                def _cell_display_length(col_letter: str, row_idx: int):
                    column_idx = column_index_from_string(col_letter)
                    cell_value = ws.cell(row=row_idx, column=column_idx).value
                    if isinstance(cell_value, str) and cell_value.startswith('='):
                        df_row_idx = row_idx - data_start_row
                        if 0 <= df_row_idx < len(df_processed_final):
                            col_name = df_col_letter_map.get(col_letter)
                            if col_name is not None:
                                df_value = df_processed_final.iloc[df_row_idx][col_name]
                                return _value_length(df_value)
                        return 0
                    return _value_length(cell_value)

                def _width_from_length(length: int):
                    if length <= 0:
                        return 8
                    return min(length + 1.5, 60)

                for col_idx in range(1, ws.max_column + 1):
                    col_letter = get_column_letter(col_idx)
                    lengths = []
                    for row_idx in range(1, header_excel_row + 1):
                        lengths.append(_cell_display_length(col_letter, row_idx))
                    if total_data_rows:
                        for row_idx in range(data_start_row, last_data_row + 1):
                            lengths.append(_cell_display_length(col_letter, row_idx))
                    max_len = max(lengths) if lengths else 0
                    ws.column_dimensions[col_letter].width = _width_from_length(max_len)

                def _set_width_from_cell(col_letter: str, row_idx: int):
                    display_len = _cell_display_length(col_letter, row_idx)
                    width_val = _width_from_length(display_len)
                    if width_val:
                        ws.column_dimensions[col_letter].width = width_val
                    return width_val

                def _set_width_from_data(col_letter: str):
                    lengths = []
                    if total_data_rows:
                        for row_idx in range(data_start_row, last_data_row + 1):
                            if ws.row_dimensions[row_idx].hidden:
                                continue
                            display_len = _cell_display_length(col_letter, row_idx)
                            if display_len:
                                lengths.append(display_len)
                    if not lengths:
                        lengths.append(_cell_display_length(col_letter, header_excel_row))
                    width_val = _width_from_length(max(lengths)) if lengths else None
                    if width_val:
                        ws.column_dimensions[col_letter].width = width_val
                    return width_val

                for col in ['X', 'Y', 'Z']:
                    _set_width_from_data(col)

                _set_width_from_cell('B', 10)
                _set_width_from_cell('C', 12)
                p_width = _set_width_from_cell('P', 12)
                if not p_width:
                    p_width = _set_width_from_cell('Q', 12)
                if p_width:
                    for col in ['N', 'O', 'P', 'Q']:
                        ws.column_dimensions[col].width = p_width

                for col in ['E', 'F', 'G', 'H', 'I', 'J', 'K']:
                    ws.column_dimensions[col].hidden = True

                last_column_letter = get_column_letter(max(ws.max_column, 26))
                auto_filter_ref = f"A{header_excel_row}:{last_column_letter}{ws.max_row if total_data_rows else header_excel_row}"
                ws.auto_filter.ref = auto_filter_ref

                def _last_row_with_value(column_letter: str, start_row: int) -> int:
                    for row_idx in range(ws.max_row, start_row - 1, -1):
                        cell_value = ws[f"{column_letter}{row_idx}"].value
                        if cell_value not in (None, ""):
                            return row_idx
                    return start_row - 1

                last_row_with_d = _last_row_with_value('D', data_start_row)
                if last_row_with_d >= data_start_row:
                    for row_idx in range(data_start_row, last_row_with_d + 1):
                        df_row_idx = row_idx - data_start_row
                        if 0 <= df_row_idx < len(df_processed_final):
                            row_series = df_processed_final.iloc[df_row_idx]
                            for col_letter, col_name in (('Y', 'Y'), ('Z', 'Z')):
                                target_cell = ws[f"{col_letter}{row_idx}"]
                                if target_cell.value in (None, ""):
                                    value = row_series.get(col_name)
                                    if not (pd.isna(value) or value == ""):
                                        target_cell.value = value

                kitid_col_index = None
                for col_idx in range(1, ws.max_column + 1):
                    if ws.cell(row=header_excel_row, column=col_idx).value == 'KITID':
                        kitid_col_index = col_idx
                        break

                if kitid_col_index is not None:
                    ws.auto_filter.add_filter_column(kitid_col_index - 1, ['KITM', 'SDRWM', '9DRWM'])

                wb.save(processed_filepath)

                df_main = df_processed_final.rename(columns={'PN':'group_key', 'DRAWING No.':'compare_val', 'K':'filename', 'QTY':'qty', 'KITID':'type'})
                
            except Exception as e:
                tb_str = traceback.format_exc()
                full_error_message = f"엑셀 사전 처리 중 오류 발생: {e}\n\n--- TRACEBACK ---\n{tb_str}"
                emit_error(f"엑셀 사전 처리 중 오류 발생: {e}", full_error_message)
                return

            emit_progress(10, "사전 처리 완료. 메인 작업 시작...", gui_log=f"사전 처리된 파일 저장: {processed_filepath}")
            
            df_main.dropna(subset=['filename', 'compare_val', 'group_key'], inplace=True)
            df_main['type'] = df_main['type'].astype(str)
            df_main['qty'] = pd.to_numeric(df_main['qty'], errors='coerce').fillna(0)

            # --- SLDPRT 처리 ---
            emit_progress(15, "SLDPRT 처리 중...", gui_log="SLDPRT 처리 시작...")
            filtered_df_sldprt = df_main[df_main['type'].isin(['KITM', 'SDRWM'])]
            numeric_filenames_k = pd.to_numeric(filtered_df_sldprt['filename'], errors='coerce')
            candidate_list_k = numeric_filenames_k.dropna().astype('Int64').astype(str).tolist()

            if not os.path.exists(SLDPRT_EXCLUSION_FILE_PATH):
                raise FileNotFoundError(f"SLDPRT 제외 목록 없음: {SLDPRT_EXCLUSION_FILE_PATH}")
            df_exclude_sldprt = pd.read_excel(SLDPRT_EXCLUSION_FILE_PATH, usecols="A", header=None, names=['filename'])
            numeric_exclude_sldprt = pd.to_numeric(df_exclude_sldprt['filename'], errors='coerce')
            exclude_set_sldprt = set(numeric_exclude_sldprt.dropna().astype('Int64').astype(str).tolist())
            files_to_find_sldprt = sorted(list(set([name for name in candidate_list_k if name not in exclude_set_sldprt])))
            
            log_for_excel(f"\n--- 최종 SLDPRT 복사 대상 ({len(files_to_find_sldprt)}개) ---")
            for item in files_to_find_sldprt: log_for_excel(item)
            log_for_excel("--------------------------------")
            
            emit_progress(25, "SLDPRT 파일 검색/복사 중...")
            all_source_files = os.listdir(source)
            copied_sldprt_pairs = []
            for base_name in files_to_find_sldprt:
                for actual_filename in all_source_files:
                    if base_name in actual_filename and actual_filename.lower().endswith(TARGET_SLDPRT_EXTENSION):
                        source_path = os.path.join(source, actual_filename)
                        dest_path = os.path.join(sldprt_dest_folder, actual_filename)
                        if not os.path.exists(dest_path):
                            shutil.copy2(source_path, dest_path)
                            copied_sldprt_pairs.append((base_name, actual_filename))

            log_for_excel(f"\n--- SLDPRT 복사 성공 ({len(copied_sldprt_pairs)}개) ---")
            for pair in copied_sldprt_pairs: log_for_excel(pair[1])
            log_for_excel("--------------------------------")
            
            emit_progress(50, "SLDPRT 결과 저장 중...", gui_log=f"SLDPRT 복사 완료: {len(copied_sldprt_pairs)}개")
            if copied_sldprt_pairs:
                b_column_files = [pair[1] for pair in copied_sldprt_pairs]
                a_column_base_names = list(dict.fromkeys([pair[0] for pair in copied_sldprt_pairs]))
                df_report = pd.DataFrame({'사용된 도면번호 (A열)': pd.Series(a_column_base_names), '복사 성공 파일 (B열)': pd.Series(b_column_files)})
                df_report.to_excel(report_excel_path, index=False)
                sldprt_final_message = f"총 {len(b_column_files)}개의 SLDPRT 파일을 복사했습니다."
            else:
                sldprt_final_message = "조건에 맞는 SLDPRT 파일이 없습니다."

            # --- JPG 처리 ---
            jpg_final_message = ""
            if copy_jpg_enabled:
                emit_progress(55, "JPG 목록 필터링 중...", gui_log="JPG 처리 시작...")
                filtered_df_jpg = df_main[df_main['type'].isin(['KITM', '9DRWM', 'SDRWM'])].copy()
                kitm_df = df_main[df_main['type'] == 'KITM'].copy()
                kitm_group_sums = kitm_df.groupby('group_key')['qty'].sum()
                bad_kitm_groups = set(kitm_group_sums[kitm_group_sums <= 0].index.astype(str))
                filtered_df_jpg = filtered_df_jpg[~((filtered_df_jpg['type'] == 'KITM') & (filtered_df_jpg['group_key'].astype(str).isin(bad_kitm_groups)))].copy()

                # ▼▼▼▼▼ 핵심 수정사항: JPG_EXCLUSION_FILE_PATH를 TOTAL_BOM_PATH로 변경 ▼▼▼▼▼
                if not os.path.exists(TOTAL_BOM_PATH): raise FileNotFoundError(f"JPG 제외 목록 없음: {TOTAL_BOM_PATH}")
                df_exclude_jpg = pd.read_excel(TOTAL_BOM_PATH, usecols="B", header=None, names=['filename'])
                # ▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲▲
                
                numeric_exclude_jpg = pd.to_numeric(df_exclude_jpg['filename'], errors='coerce')
                exclude_set_jpg = set(numeric_exclude_jpg.dropna().astype('Int64').astype(str).tolist())
                
                filtered_df_jpg.loc[:, 'compare_val_numeric'] = pd.to_numeric(filtered_df_jpg['compare_val'], errors='coerce')
                compare_numeric = filtered_df_jpg['compare_val_numeric']
                valid_compare = compare_numeric.notna()
                compare_as_str = compare_numeric[valid_compare].astype('Int64').astype(str)
                exclude_index = compare_as_str[compare_as_str.isin(exclude_set_jpg)].index

                if 'X' in filtered_df_jpg.columns:
                    exclude_index = exclude_index.union(filtered_df_jpg[filtered_df_jpg['X'].notna()].index)

                non_duplicate_df = filtered_df_jpg.drop(index=exclude_index)
                numeric_filenames_jpg = pd.to_numeric(non_duplicate_df['filename'], errors='coerce')
                files_to_find_jpg = sorted(list(set(numeric_filenames_jpg.dropna().astype('Int64').astype(str).tolist())))
                
                log_for_excel(f"\n--- 최종 JPG 복사 대상 ({len(files_to_find_jpg)}개) ---")
                for item in files_to_find_jpg: log_for_excel(item)
                log_for_excel("--------------------------------")

                emit_progress(70, "JPG 폴더 스캔 중...")
                if not os.path.exists(JPG_SOURCE_PATH): raise FileNotFoundError(f"고정된 JPG 원본 폴더 없음: {JPG_SOURCE_PATH}")
                jpg_file_map = {}
                for dirpath, _, filenames in os.walk(JPG_SOURCE_PATH):
                    for f in filenames:
                        if f.lower().endswith(TARGET_JPG_EXTENSION):
                            jpg_file_map[f] = os.path.join(dirpath, f)
                
                emit_progress(85, "JPG 파일을 검색/복사 중...")
                os.makedirs(jpg_dest_folder, exist_ok=True)
                
                copied_jpg_files = []
                for base_name in files_to_find_jpg:
                    pattern = re.compile(f"{re.escape(base_name)}(?![0-9])")
                    for jpg_filename, full_path in jpg_file_map.items():
                        if pattern.search(jpg_filename):
                            dest_path = os.path.join(jpg_dest_folder, jpg_filename)
                            if not os.path.exists(dest_path):
                                shutil.copy2(full_path, dest_path)
                                copied_jpg_files.append(jpg_filename)
                
                log_for_excel(f"\n--- JPG 복사 성공 ({len(copied_jpg_files)}개) ---")
                for item in copied_jpg_files: log_for_excel(item)
                log_for_excel("--------------------------------")
                
                jpg_final_message = f"\n추가로 {len(copied_jpg_files)}개의 JPG 파일을 복사했습니다."
                emit_progress(95, "JPG 처리 완료", gui_log=f"JPG 복사 완료: {len(copied_jpg_files)}개")

            # ==== 최종 마무리 ====
            final_message = (sldprt_final_message + jpg_final_message).strip()
            log_for_excel(f"\n========== 작업 완료: {dt.datetime.now().strftime('%Y-%m-%d %H:%M%S')} ==========")
            
            try:
                pd.DataFrame(excel_log_records, columns=["Log"]).to_excel(log_excel_path, index=False, sheet_name='LOG')
            except Exception as e:
                final_message += f"\n로그 저장 실패: {e}"
            
            emit_done('✅ 작업 완료!', final_message, log_save_path=log_excel_path)

        except Exception as e:
            tb_str = traceback.format_exc()
            full_error_message = f"작업 중 예외 발생: {e}\n\n--- TRACEBACK ---\n{tb_str}"
            log_for_excel(full_error_message)
            emit_error(f"오류 발생: {e}", full_error_message)

if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()
